"""
HR Report generator — 6-sheet Excel workbook using openpyxl.

Sheet 1: Summary              — totals + quick leaderboard
Sheet 2: Ranked Candidates    — per-concept detail for passing candidates
Sheet 3: Rejected Candidates  — rejection reasons + HR note
Sheet 4: Per-Question Detail  — one row per (candidate, concept)
Sheet 5: Keyword Gap Report   — keywords no candidate covered
Sheet 6: How to Read          — plain-language guide for HR

Colour coding:
  Green  (#C6EFCE) → ✅ Human Review
  Yellow (#FFEB9C) → 🟡 Shortlist
  Red    (#FFC7CE) → ❌ Auto-Reject
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

_GREEN = PatternFill("solid", fgColor="C6EFCE")
_YELLOW = PatternFill("solid", fgColor="FFEB9C")
_RED = PatternFill("solid", fgColor="FFC7CE")
_BOLD = Font(bold=True)


def _fill(segment: str) -> Optional[PatternFill]:
    if "Human Review" in segment:
        return _GREEN
    if "Shortlist" in segment:
        return _YELLOW
    if "Auto-Reject" in segment:
        return _RED
    return None


def _colour_row(ws, fill: Optional[PatternFill]) -> None:
    if fill:
        for cell in ws[ws.max_row]:
            cell.fill = fill


def _auto_width(ws) -> None:
    for col in ws.columns:
        width = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 4, 70)


def _bold_row(ws, row_idx: int) -> None:
    for cell in ws[row_idx]:
        cell.font = _BOLD


# ─────────────────────────────────────────────────────────────────────────────
# Sheet builders
# ─────────────────────────────────────────────────────────────────────────────

def _sheet_summary(wb: Workbook, leaderboard: list[dict]) -> None:
    ws = wb.create_sheet("Summary")
    passing = [e for e in leaderboard if not e["reject"]]
    rejected = [e for e in leaderboard if e["reject"]]

    ws.append(["Assessment Leaderboard — Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Metric", "Value"])
    _bold_row(ws, ws.max_row)
    ws.append(["Total Assessed", len(leaderboard)])
    ws.append(["Total Passed", len(passing)])
    ws.append(["Total Rejected", len(rejected)])
    ws.append([])

    ws.append(["Rank", "Name", "Final Score", "Avg Depth", "Segment", "Rejected", "Reject Reason"])
    _bold_row(ws, ws.max_row)

    for entry in leaderboard:
        ws.append([
            entry.get("rank") or "—",
            entry["name"],
            f"{entry['final_score']:.4f}",
            f"{entry['avg_depth']:.4f}",
            entry["segment"],
            "Yes" if entry["reject"] else "No",
            entry.get("reject_reason") or "",
        ])
        _colour_row(ws, _fill(entry["segment"]))

    _auto_width(ws)


def _knowledge_source(q: dict) -> str:
    return "AI General Knowledge" if q.get("is_external") else "Internal DB"


def _sheet_ranked(wb: Workbook, leaderboard: list[dict]) -> None:
    ws = wb.create_sheet("Ranked Candidates")
    passing = [e for e in leaderboard if not e["reject"]]

    headers = [
        "Rank", "Name", "Final Score", "Avg Depth", "Segment",
        "Concept", "Knowledge Source", "Coverage Score", "Depth Score", "Comparative Score",
        "Matched Keywords", "Missing Keywords", "Recovered Keywords",
    ]
    ws.append(headers)
    _bold_row(ws, 1)

    for entry in passing:
        fill = _fill(entry["segment"])
        first = True
        for cname, q in entry.get("per_question", {}).items():
            row = []
            if first:
                row += [
                    entry["rank"],
                    entry["name"],
                    f"{entry['final_score']:.4f}",
                    f"{entry['avg_depth']:.4f}",
                    entry["segment"],
                ]
                first = False
            else:
                row += ["", "", "", "", ""]
            row += [
                cname,
                _knowledge_source(q),
                f"{q['coverage_score']:.4f}",
                f"{q['depth_score']:.4f}",
                f"{q['comparative_score']:.4f}",
                ", ".join(q.get("matched_keywords", [])),
                ", ".join(q.get("missing_keywords", [])),
                ", ".join(q.get("semantic_recovered", [])),
            ]
            ws.append(row)
            _colour_row(ws, fill)

    _auto_width(ws)


def _sheet_rejected(wb: Workbook, leaderboard: list[dict]) -> None:
    ws = wb.create_sheet("Rejected Candidates")
    rejected = [e for e in leaderboard if e["reject"]]

    ws.append(["HR Note: These candidates scored below threshold on both coverage and depth."])
    ws["A1"].font = Font(bold=True, italic=True)
    ws.append([])
    ws.append(["Name", "Final Score", "Avg Depth", "Reject Reason"])
    _bold_row(ws, ws.max_row)

    for entry in rejected:
        ws.append([
            entry["name"],
            f"{entry['final_score']:.4f}",
            f"{entry['avg_depth']:.4f}",
            entry.get("reject_reason") or "",
        ])
        _colour_row(ws, _RED)

    _auto_width(ws)


def _sheet_per_question(wb: Workbook, leaderboard: list[dict]) -> None:
    ws = wb.create_sheet("Per-Question Detail")
    headers = [
        "Name", "Concept", "Knowledge Source", "Coverage Score", "Depth Score",
        "Comparative Score", "Matched Keywords", "Missing Keywords", "Recovered Keywords",
    ]
    ws.append(headers)
    _bold_row(ws, 1)

    for entry in leaderboard:
        fill = _fill(entry["segment"])
        for cname, q in entry.get("per_question", {}).items():
            ws.append([
                entry["name"],
                cname,
                _knowledge_source(q),
                f"{q['coverage_score']:.4f}",
                f"{q['depth_score']:.4f}",
                f"{q['comparative_score']:.4f}",
                ", ".join(q.get("matched_keywords", [])),
                ", ".join(q.get("missing_keywords", [])),
                ", ".join(q.get("semantic_recovered", [])),
            ])
            _colour_row(ws, fill)

    _auto_width(ws)


def _sheet_keyword_gap(wb: Workbook, keyword_gap_report: dict) -> None:
    ws = wb.create_sheet("Keyword Gap Report")
    ws.append(["Concept", "Knowledge Source", "Universal Gaps (keywords no candidate covered)"])
    _bold_row(ws, 1)

    for cname, gaps in keyword_gap_report.items():
        is_ext = "[External Knowledge]" in cname
        source = "AI General Knowledge" if is_ext else "Internal DB"
        if gaps == "all covered":
            ws.append([cname, source, "all covered"])
        else:
            ws.append([cname, source, ", ".join(gaps) if isinstance(gaps, list) else str(gaps)])
        if gaps != "all covered" and gaps:
            _colour_row(ws, _YELLOW)

    _auto_width(ws)


def _sheet_guide(wb: Workbook) -> None:
    ws = wb.create_sheet("How to Read This Report")

    rows = [
        ["How to Read This Report"],
        [],
        ["SEGMENTS"],
        ["✅ Human Review", "Score ≥ 0.75. Strong candidate — prioritise for scheduling."],
        ["🟡 Shortlist",   "Score 0.50–0.74. Acceptable candidate — review at your discretion."],
        ["❌ Auto-Reject", "Score < 0.50 with insufficient depth. Unlikely fit."],
        [],
        ["SCORES EXPLAINED"],
        ["Final Score",
         "Weighted average across all concepts: (coverage_norm × 0.60) + (depth × 0.40). "
         "Normalised against the strongest candidate in the pool — scores are RELATIVE."],
        ["Coverage Score",
         "Fraction of required technical keywords present in the answer (exact or semantic match)."],
        ["Depth Score",
         "GPT-4o-mini quality rating from 0.0 (empty/wrong) to 1.0 (production-grade)."],
        ["Comparative Score",
         "Per-concept weighted score: coverage_norm × 0.60 + depth × 0.40."],
        [],
        ["DEPTH COMPENSATION"],
        ["",
         "A candidate with final_score < 0.50 but avg_depth > 0.50 is NOT auto-rejected. "
         "Deep explanation compensates for keyword gaps — they appear in 🟡 Shortlist."],
        [],
        ["TIEBREAKING"],
        ["",
         "When two candidates are within 0.05 of each other in final score, "
         "GPT-4o runs a best-of-3 deterministic comparison to determine rank order."],
        [],
        ["KNOWLEDGE SOURCE"],
        ["Internal DB",
         "Questions sourced from the internal Knowledge DB. Required keywords are "
         "concept-level ground truth — highly reliable."],
        ["AI General Knowledge",
         "This tool was not in the internal Knowledge DB. Questions were generated "
         "and graded using AI general knowledge. Results are indicative but not "
         "backed by curated concept metadata. Consider adding this tool to the KB."],
        [],
        ["OVERRIDE POLICY"],
        ["",
         "Scores are AI-generated signals, not final verdicts. HR may override any "
         "segment decision based on portfolio, references, or live interview performance."],
    ]

    for row_data in rows:
        ws.append(row_data)

    ws["A1"].font = Font(bold=True, size=13)
    for row_idx in (3, 8, 14, 17, 20):
        try:
            _bold_row(ws, row_idx)
        except Exception:
            pass

    _auto_width(ws)


# ─────────────────────────────────────────────────────────────────────────────
# Public entry point
# ─────────────────────────────────────────────────────────────────────────────

def generate_hr_report(
    leaderboard: list[dict],
    keyword_gap_report: dict,
    jr_id: int,
    output_dir: Path,
) -> Path:
    """
    Generates the 6-sheet HR Excel workbook and saves it to output_dir.
    Returns the full path of the saved file.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"hr_report_jr{jr_id}_{ts}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    _sheet_summary(wb, leaderboard)
    _sheet_ranked(wb, leaderboard)
    _sheet_rejected(wb, leaderboard)
    _sheet_per_question(wb, leaderboard)
    _sheet_keyword_gap(wb, keyword_gap_report)
    _sheet_guide(wb)

    wb.save(path)
    logger.info("[hr_report] Saved → %s", path)
    return path
